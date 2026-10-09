"""Checks exported PowerPoint files with an independent reader (python-pptx) and, when the
ECMA-376 schemas are given, validates the slide XML (markup-compatibility blocks are replaced
by their fallback, as a consumer without the extension would do).

usage: pptx_check.py [--xsd DIR] file.pptx...
"""
import argparse
import os
import re
import sys
import zipfile

from pptx import Presentation  # python-pptx

MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def check_xsd(path, schema):
    from lxml import etree

    errors = []
    with zipfile.ZipFile(path) as z:
        for name in sorted(n for n in z.namelist() if re.match(r"ppt/(slides/slide\d+|presentation)\.xml$", n)):
            doc = etree.fromstring(z.read(name))
            for ac in list(doc.iter("{%s}AlternateContent" % MC)):
                fallback = ac.find("{%s}Fallback" % MC)
                parent = ac.getparent()
                index = parent.index(ac)
                for k, child in enumerate(list(fallback) if fallback is not None else []):
                    parent.insert(index + k, child)
                parent.remove(ac)
            if not schema.validate(doc):
                errors.append(f"{name}: {schema.error_log.last_error}")
    return errors


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xsd", help="folder with pml.xsd and the schemas it imports")
    ap.add_argument("files", nargs="+")
    args = ap.parse_args()
    schema = None
    if args.xsd:
        from lxml import etree

        schema = etree.XMLSchema(etree.parse(os.path.join(args.xsd, "pml.xsd")))
    failed = False
    for f in args.files:
        p = Presentation(f)
        shapes = sum(len(s.shapes) for s in p.slides)
        print(f"{os.path.basename(f)}: {len(p.slides)} slide(s), {shapes} shape(s)")
        if len(p.slides) == 0 or shapes == 0:
            print("  empty presentation")
            failed = True
        if schema is not None:
            for e in check_xsd(f, schema):
                print("  XSD:", e)
                failed = True
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
